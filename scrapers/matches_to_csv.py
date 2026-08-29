"""
Reads every responses_matches/matches_series*.json file (from
fetch_get_matches_by_series.py), dedupes matches by matchId, and writes
a single formatted Excel report in the "Regional Matches" template style
(NO. / DATE / DAY / REGION / NAME OF TOURNAMENT / TEAM 1 / TEAM 2 / VENUE
/ REMARKS, with a merged month header, bold column headers, borders, and
auto-sized columns).

This replaces the old matches_to_csv.py flat-CSV export with the same
loading/dedup logic, but reshapes the data into the report layout used
by the "44. Hamisi Work" template.

Usage:
    python matches_to_excel.py
"""

import glob
import json
import os

import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RESPONSES_DIR = "responses_matches"

# Where the formatted .xlsx should be written.
EXPORT_FOLDER = "D:\\OneDrive\\1. Project\\TCA Analyst\\1. Projects\\44. Hamisi Work"

# Mapping of keyword -> region label to search for in tournament name /
# venue. Some keywords (Azania, Kitunda, Kigamboni, Kawe) are areas within
# Dar es Salaam, so they map to 'Dar' rather than being their own region.
REGION_KEYWORDS = {
    'MOROGORO': 'Morogoro',
    'GAIRO': 'Gairo',
    'GEITA': 'Geita',
    'ARUSHA': 'Arusha',
    'DODOMA': 'Dodoma',
    'KONGWA': 'Kongwa',
    'AZANIA': 'Dar',
    'BAGAMOYO': 'Bagamoyo',
    'KITUNDA': 'Dar',
    'KIGAMBONI': 'Dar',
    'TANGA': 'Tanga',
    'KILINDI': 'Kilindi',
    'KISARAWE': 'Kisarawe',
    'KILIMANJARO': 'Kilimanjaro',
    'KAWE': 'Dar',
    'CHANIKA': 'Chanika',
    'CHAMWINO': 'Chamwino',
    'MWANZA': 'Mwanza',
    'ZANZIBAR': 'Zanzibar',
    'IRINGA': 'Iringa',
    'PANDAMBILI': 'Pandambili',
    'KOROGWE': 'Korogwe',
}


def load_all_matches():
    """Same load/dedupe logic as matches_to_csv.py: later files win on a
    matchId conflict, since a given match should be identical regardless
    of which series query surfaced it."""
    files = sorted(glob.glob(os.path.join(RESPONSES_DIR, "matches_series*.json")))
    print(f"Found {len(files)} series batch files in {RESPONSES_DIR}/")

    matches_by_id = {}
    for path in files:
        with open(path, "r", encoding="utf-8") as f:
            raw = json.load(f)
        for match in raw.get("data", []):
            match_id = match.get("matchId")
            if match_id is None:
                continue
            matches_by_id[match_id] = match

    return matches_by_id


def assign_region(venue, tournament_name):
    venue_lower = str(venue).lower()

    # 1. Existing specific venue -> region rules take priority
    if any(x in venue_lower for x in ['annadil burhani', 'leaders', 'gymkhana', 'udsm']):
        return 'Dar'
    elif 'sua ground' in venue_lower:
        return 'Morogoro'
    elif 'usagara' in venue_lower:
        return 'Tanga'
    elif 'dodoma jiji' in venue_lower:
        return 'Dodoma'
    elif 'mwanakalenge' in venue_lower:
        return 'Bagamoyo'
    elif 'gairo' in venue_lower:
        return 'Gairo'
    elif 'butimba' in venue_lower:
        return 'Mwanza'

    # 2. Fall back to scanning the tournament name for a known region keyword
    tournament_upper = str(tournament_name).upper()
    for keyword, region in REGION_KEYWORDS.items():
        if keyword in tournament_upper:
            return region

    # 3. Fall back to scanning the venue text itself for a known region keyword
    venue_upper = str(venue).upper()
    for keyword, region in REGION_KEYWORDS.items():
        if keyword in venue_upper:
            return region

    return ' '


def build_report_dataframe(matches_by_id):
    """Reshape raw match dicts (matchId, teamOneName, teamTwoName,
    matchDate, seriesName, location, ...) into the report columns."""
    rows = []
    for match_id in sorted(matches_by_id.keys()):
        m = matches_by_id[match_id]
        rows.append({
            'dates': m.get('matchDate'),
            # Prefer seriesName (the tournament), fall back to clubName
            'event_name': m.get('seriesName') or m.get('clubName') or '',
            'venue': m.get('location') or '',
            'first_bat': m.get('teamOneName') or m.get('teamOneCode') or '',
            'second_bat': m.get('teamTwoName') or m.get('teamTwoCode') or '',
        })

    new_df = pd.DataFrame(rows)

    # Drop rows with no usable date — they can't be placed in a monthly report
    new_df = new_df.dropna(subset=['dates'])
    new_df = new_df[new_df['dates'] != '']

    cleaned_df = new_df[['dates', 'event_name', 'venue', 'first_bat', 'second_bat']].copy()
    cleaned_df.columns = ['DATE', 'NAME OF TOURNAMENT', 'VENUE', 'TEAM 1', 'TEAM 2']

    cleaned_df['DATE'] = pd.to_datetime(cleaned_df['DATE'])
    cleaned_df = cleaned_df.sort_values('DATE').reset_index(drop=True)

    # Use the last available match date to determine which month/year this
    # report covers, then keep only matches that fall within that month.
    last_date = cleaned_df['DATE'].max()
    target_year, target_month = last_date.year, last_date.month
    cleaned_df = cleaned_df[
        (cleaned_df['DATE'].dt.year == target_year) &
        (cleaned_df['DATE'].dt.month == target_month)
    ].reset_index(drop=True)

    cleaned_df.insert(0, 'NO.', range(1, len(cleaned_df) + 1))
    cleaned_df['DAY'] = cleaned_df['DATE'].dt.day_name()

    cleaned_df['REGION'] = cleaned_df.apply(
        lambda row: assign_region(row['VENUE'], row['NAME OF TOURNAMENT']), axis=1
    )

    cleaned_df['REMARKS'] = ''

    cleaned_df = cleaned_df[[
        'NO.', 'DATE', 'DAY', 'REGION', 'NAME OF TOURNAMENT',
        'TEAM 1', 'TEAM 2', 'VENUE', 'REMARKS'
    ]]

    return cleaned_df


def write_formatted_excel(cleaned_df):
    if cleaned_df.empty:
        print("No dated matches to write — skipping Excel export.")
        return

    first_date = cleaned_df['DATE'].iloc[0]
    month_name_upper = first_date.strftime('%B').upper()
    month_name_proper = first_date.strftime('%B')
    year = first_date.strftime('%Y')

    os.makedirs(EXPORT_FOLDER, exist_ok=True)
    output_file = os.path.join(EXPORT_FOLDER, f"{month_name_proper} {year} Regional Matches.xlsx")

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # Leave row 1 for the merged header, row 2 blank, table from row 3
        cleaned_df.to_excel(writer, sheet_name='Cricket Matches', index=False, startrow=2)

        worksheet = writer.sheets['Cricket Matches']

        num_columns = len(cleaned_df.columns)
        last_column_letter = worksheet.cell(row=3, column=num_columns).column_letter

        worksheet.merge_cells(f'A1:{last_column_letter}1')
        worksheet['A1'] = f'REGIONAL MATCHES FOR THE MONTH OF {month_name_upper} {year}'
        worksheet['A1'].font = Font(bold=True)
        worksheet['A1'].alignment = Alignment(horizontal='left', vertical='center')

        bold_font = Font(bold=True)
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Column headers (row 3)
        for cell in worksheet[3]:
            cell.font = bold_font
            cell.alignment = center_alignment
            cell.border = thin_border

        # Data rows
        for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = center_alignment
                cell.border = thin_border

        # Auto-size columns
        for col_idx in range(1, num_columns + 1):
            column_letter = get_column_letter(col_idx)
            max_length = 0
            for row in range(3, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)

        # Consistent row height
        for row in range(3, worksheet.max_row + 1):
            worksheet.row_dimensions[row].height = 15

        # DATE column formatting
        for row in range(4, worksheet.max_row + 1):
            cell = worksheet[f'B{row}']
            cell.number_format = 'DD-MMM-YYYY'
            cell.border = thin_border

    print(f"Excel file '{output_file}' created successfully with formatting!")
    print(f"Total rows: {len(cleaned_df)}")


def main():
    matches_by_id = load_all_matches()
    if not matches_by_id:
        print("No matches found — check that fetch_get_matches_by_series.py has run "
              f"and populated {RESPONSES_DIR}/")
        return

    print(f"Total unique matches: {len(matches_by_id)}")

    cleaned_df = build_report_dataframe(matches_by_id)
    write_formatted_excel(cleaned_df)


if __name__ == "__main__":
    main()